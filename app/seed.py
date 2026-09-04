import argparse
from pathlib import Path
from typing import Optional, Union

import openpyxl

from app.db import (
    capture_snapshot,
    connect,
    insert_expense,
    insert_income_source,
    replace_accounts,
)
from app.projections import Account, Expense, IncomeSource, default_bank


def _to_pence(value, multiplier=100):
    pence = round(float(value or 0) * multiplier)
    return pence


def _find_header_row(ws, text, column=1):
    result = None
    for row in range(1, 100):
        value = ws.cell(row=row, column=column).value
        if value and str(value).upper() == text.upper():
            result = row
            break
    return result


def _read_savings(ws) -> list[Account]:
    accounts: list[Account] = []
    for row in range(3, 28):
        name = ws.cell(row=row, column=1).value
        if not name:
            break
        accounts.append(
            Account(
                name=str(name),
                balance_pence=_to_pence(ws.cell(row=row, column=2).value),
                annual_rate_bp=_to_pence(ws.cell(row=row, column=3).value),
                monthly_allocation_pence=_to_pence(ws.cell(row=row, column=4).value),
                bank=default_bank(str(name)),
            )
        )
    return accounts


def _read_credit_cards(ws) -> list[Account]:
    accounts: list[Account] = []
    start = _find_header_row(ws, "CREDIT CARDS")
    if start is not None:
        for row in range(start + 2, start + 30):
            name = ws.cell(row=row, column=1).value
            if not name or "TOTAL" in str(name).upper():
                break
            accounts.append(
                Account(
                    name=str(name),
                    balance_pence=_to_pence(ws.cell(row=row, column=2).value),
                    kind="credit_card",
                )
            )
    return accounts


def _read_table(ws, header_text, header_col, name_col, value_col):
    rows: list[tuple[str, int]] = []
    start = _find_header_row(ws, header_text, column=header_col)
    if start is not None:
        for row in range(start + 2, start + 30):
            name = ws.cell(row=row, column=name_col).value
            if not name or "TOTAL" in str(name).upper():
                break
            amount = ws.cell(row=row, column=value_col).value
            rows.append((str(name), _to_pence(amount)))
    return rows


def read_accounts_from_xlsx(path: Union[str, Path]) -> list[Account]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Inputs"]
    accounts = _read_savings(worksheet) + _read_credit_cards(worksheet)
    return accounts


def read_income_from_xlsx(path: Union[str, Path]) -> list[IncomeSource]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Inputs"]
    rows = _read_table(worksheet, "MONTHLY INCOME", 6, 6, 7)
    sources = [
        IncomeSource(name=name, monthly_amount_pence=amount) for name, amount in rows
    ]
    return sources


def read_expenses_from_xlsx(path: Union[str, Path]) -> list[Expense]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Inputs"]
    monthly_rows = _read_table(worksheet, "MONTHLY EXPENSES", 6, 6, 7)
    yearly_rows = _read_table(worksheet, "YEARLY EXPENSES", 9, 9, 10)
    monthly = [
        Expense(name=name, amount_pence=amount, cadence="monthly")
        for name, amount in monthly_rows
    ]
    yearly = [
        Expense(name=name, amount_pence=amount, cadence="yearly")
        for name, amount in yearly_rows
    ]
    return monthly + yearly


def seed_database(
    xlsx_path: Union[str, Path], db_path: Union[str, Path]
) -> dict[str, int]:
    accounts = read_accounts_from_xlsx(xlsx_path)
    income = read_income_from_xlsx(xlsx_path)
    expenses = read_expenses_from_xlsx(xlsx_path)
    conn = connect(db_path)
    capture_snapshot(conn, trigger="Manual", label="Before seed")
    with conn:
        conn.execute("DELETE FROM income_sources")
        conn.execute("DELETE FROM expenses")
    replace_accounts(conn, accounts)
    for source in income:
        insert_income_source(conn, source)
    for expense in expenses:
        insert_expense(conn, expense)
    return {
        "accounts": len(accounts),
        "income_sources": len(income),
        "expenses": len(expenses),
    }


def main(argv: Optional[list[str]] = None) -> None:
    parser = argparse.ArgumentParser(
        description="Seed the database from a Finances.xlsx file"
    )
    parser.add_argument(
        "xlsx",
        type=Path,
        nargs="?",
        default=Path("data/Finances.xlsx"),
        help="Path to xlsx (default: data/Finances.xlsx)",
    )
    parser.add_argument(
        "--db",
        type=Path,
        default=Path("data/finance.db"),
        help="Path to SQLite database (default: data/finance.db)",
    )
    args = parser.parse_args(argv)
    args.db.parent.mkdir(parents=True, exist_ok=True)
    counts = seed_database(args.xlsx, args.db)
    summary = ", ".join(f"{count} {name}" for name, count in counts.items())
    print(f"Seeded {summary} from {args.xlsx} into {args.db}")


if __name__ == "__main__":
    main()
