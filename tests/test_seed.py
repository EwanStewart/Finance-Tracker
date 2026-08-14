import openpyxl

from app.db import (
    connect,
    insert_account,
    list_accounts,
    list_expenses,
    list_income_sources,
    list_snapshots,
)
from app.projections import Account
from app.seed import (
    read_accounts_from_xlsx,
    read_expenses_from_xlsx,
    read_income_from_xlsx,
    seed_database,
)


def _write_inputs_sheet(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "SAVINGS ACCOUNTS"
    ws["A2"] = "Account Name"
    ws["B2"] = "Balance (£)"
    ws["C2"] = "Annual Rate (%)"
    ws["D2"] = "Monthly Allocation (£)"
    ws["A3"] = "Test Savings"
    ws["B3"] = 1000.50
    ws["C3"] = 3.5
    ws["D3"] = 100.0
    ws["A4"] = "Other Savings"
    ws["B4"] = 200.0
    ws["C4"] = 5.0
    ws["D4"] = 0
    ws["A28"] = "TOTAL SAVINGS"
    ws["A30"] = "CREDIT CARDS"
    ws["A31"] = "Account Name"
    ws["B31"] = "Balance (£)"
    ws["A32"] = "Test Card"
    ws["B32"] = -250.0
    ws["A35"] = "TOTAL CC BALANCE"
    ws["F1"] = "MONTHLY INCOME"
    ws["F2"] = "Source"
    ws["G2"] = "Amount (£)"
    ws["F3"] = "Salary"
    ws["G3"] = 2954.0
    ws["F4"] = "TOTAL INCOME"
    ws["F15"] = "MONTHLY EXPENSES"
    ws["F16"] = "Category"
    ws["G16"] = "Amount (£)"
    ws["F17"] = "Rent"
    ws["G17"] = 200.0
    ws["F18"] = "Phone"
    ws["G18"] = 10.9
    ws["F37"] = "TOTAL EXPENSES"
    ws["I15"] = "YEARLY EXPENSES"
    ws["I16"] = "Category"
    ws["J16"] = "Amount (£)"
    ws["I17"] = "Car Insurance"
    ws["J17"] = 700.0
    ws["I18"] = "Road Tax"
    ws["J18"] = 200.0
    ws["I37"] = "TOTAL YEARLY"
    wb.save(path)


def test_read_accounts_parses_savings_and_credit_cards(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)

    accounts = read_accounts_from_xlsx(xlsx)

    assert accounts == [
        Account(
            name="Test Savings",
            balance_pence=100_050,
            annual_rate_bp=350,
            monthly_allocation_pence=10_000,
        ),
        Account(
            name="Other Savings",
            balance_pence=20_000,
            annual_rate_bp=500,
            monthly_allocation_pence=0,
        ),
        Account(
            name="Test Card",
            balance_pence=-25_000,
            annual_rate_bp=0,
            monthly_allocation_pence=0,
            kind="credit_card",
        ),
    ]


def test_seed_database_replaces_accounts(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)
    db_path = tmp_path / "test.db"

    counts = seed_database(xlsx, db_path)

    assert counts["accounts"] == 3
    conn = connect(db_path)
    assert len(list_accounts(conn)) == 3


def test_seed_database_snapshots_the_old_data_before_replacing_it(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)
    db_path = tmp_path / "test.db"
    insert_account(connect(db_path), Account(name="Old account", balance_pence=42))

    seed_database(xlsx, db_path)

    snapshots = list_snapshots(connect(db_path))
    assert snapshots[0].label == "Before seed"
    assert snapshots[0].payload["accounts"][0]["name"] == "Old account"


def test_read_income_parses_sources_until_total_row(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)

    sources = read_income_from_xlsx(xlsx)

    assert len(sources) == 1
    assert sources[0].name == "Salary"
    assert sources[0].monthly_amount_pence == 2_954_00


def test_read_expenses_parses_monthly_and_yearly(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)

    expenses = read_expenses_from_xlsx(xlsx)

    monthly = [e for e in expenses if e.cadence == "monthly"]
    yearly = [e for e in expenses if e.cadence == "yearly"]
    assert {(e.name, e.amount_pence) for e in monthly} == {
        ("Rent", 200_00),
        ("Phone", 10_90),
    }
    assert {(e.name, e.amount_pence) for e in yearly} == {
        ("Car Insurance", 700_00),
        ("Road Tax", 200_00),
    }


def test_seeded_credit_cards_have_kind_credit_card(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)
    db_path = tmp_path / "test.db"
    seed_database(xlsx, db_path)

    conn = connect(db_path)
    accounts = list_accounts(conn)

    by_kind = {a.kind for a in accounts}
    assert by_kind == {"savings", "credit_card"}
    assert all(a.kind == "credit_card" for a in accounts if a.balance_pence < 0)


def test_seed_database_populates_income_and_expenses(tmp_path):
    xlsx = tmp_path / "fixture.xlsx"
    _write_inputs_sheet(xlsx)
    db_path = tmp_path / "test.db"

    counts = seed_database(xlsx, db_path)

    assert counts == {"accounts": 3, "income_sources": 1, "expenses": 4}
    conn = connect(db_path)
    assert len(list_income_sources(conn)) == 1
    assert len(list_expenses(conn)) == 4
